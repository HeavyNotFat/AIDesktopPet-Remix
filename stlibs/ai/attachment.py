from __future__ import annotations

import base64
import binascii
import mimetypes
import os

IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml",
    ".toml", ".ini", ".cfg", ".conf", ".env", ".py", ".js", ".ts", ".tsx", ".jsx", ".css", ".html",
    ".htm", ".xml", ".sql", ".sh", ".bat", ".ps1", ".java", ".c", ".h", ".cpp", ".go", ".rs",
}
DOC_SUFFIXES = {".docx"} | TEXT_SUFFIXES

MAX_IMAGE_BYTES = 8 * 1024 * 1024
MAX_FILE_BYTES = 4 * 1024 * 1024
MAX_DOC_CHARS = 8000
MAX_ATTACHMENTS = 6


def encode_image(image_path):
    with open(image_path, "rb") as image_file:
        return base64.b64encode(image_file.read()).decode("utf-8")


def human_size(size) -> str:
    try:
        size = float(size)
    except (TypeError, ValueError):
        return "?"

    for unit in ("B", "KB", "MB"):
        if size < 1024 or unit == "MB":
            return f"{size:.0f} {unit}" if unit == "B" else f"{size:.1f} {unit}"
        size /= 1024
    return f"{size:.1f} MB"


def suffix_of(name) -> str:
    return os.path.splitext(str(name or ""))[1].lower()


def kind_of(name, mime="") -> str:
    if str(mime or "").startswith("image/"):
        return "image"
    if suffix_of(name) in IMAGE_SUFFIXES:
        return "image"
    return "file"


def is_readable_text(name, mime="") -> bool:
    if str(mime or "").startswith("text/"):
        return True
    if str(mime or "") in {"application/json", "application/xml", "application/x-yaml"}:
        return True
    return suffix_of(name) in TEXT_SUFFIXES


def extract_docx(data: bytes) -> str:
    """装了 python-docx 就抽正文，没装就返回空串（不影响其它功能）。"""
    try:
        import io

        import docx
    except ImportError:
        return ""

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:  # noqa: BLE001 - 坏文件不该把上传流程带崩
        return ""

    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return ""


def build(name, data: bytes, mime="") -> dict:
    """把一段字节变成附件字典；超大、读不出内容都会在 note 里说明。"""
    name = os.path.basename(str(name or "")).strip() or "未命名"
    mime = str(mime or "") or (mimetypes.guess_type(name)[0] or "")
    kind = kind_of(name, mime)
    size = len(data or b"")
    limit = MAX_IMAGE_BYTES if kind == "image" else MAX_FILE_BYTES

    attachment = {
        "kind": kind,
        "name": name,
        "mime": mime,
        "size": size,
        "text": "",
        "note": "",
    }

    if size == 0:
        attachment["note"] = "空文件"
        return attachment
    if size > limit:
        attachment["note"] = f"超过 {human_size(limit)}，没有读取内容"
        return attachment

    if kind == "image":
        attachment["data"] = base64.b64encode(data).decode("ascii")
        attachment["note"] = "图片"
        return attachment

    text = ""
    if suffix_of(name) == ".docx":
        text = extract_docx(data)
    elif is_readable_text(name, mime):
        text = decode_text(data)

    if text.strip():
        attachment["text"] = text[:MAX_DOC_CHARS]
        attachment["truncated"] = len(text) > MAX_DOC_CHARS
        attachment["note"] = "已读入正文"
    else:
        attachment["note"] = "无法读取正文（只带上了文件名）"
    return attachment


def from_path(path) -> dict:
    with open(path, "rb") as handle:
        data = handle.read()
    return build(os.path.basename(path), data, mimetypes.guess_type(str(path))[0] or "")


def from_bytes(data: bytes, name: str, mime="") -> dict:
    return build(name, data, mime)


def from_base64(data: str, name: str, mime="") -> dict:
    """网页端上传上来的就是 base64。"""
    payload = str(data or "")
    if payload.startswith("data:"):
        _head, _sep, payload = payload.partition(",")
    try:
        raw = base64.b64decode(payload, validate=False)
    except (binascii.Error, ValueError):
        raw = b""
    return build(name, raw, mime)


def normalize(items) -> list:
    """接受字典或路径，统一成附件列表（去重、限量、跳过坏数据）。"""
    result = []
    seen = set()
    for item in items or []:
        if len(result) >= MAX_ATTACHMENTS:
            break

        if isinstance(item, str):
            attachment = from_path(item) if os.path.isfile(item) else None
        elif isinstance(item, dict):
            if item.get("kind") and "size" in item and "name" in item:
                attachment = dict(item)
            elif item.get("data"):
                attachment = from_base64(item.get("data"), item.get("name", ""), item.get("mime", ""))
            elif item.get("path"):
                attachment = from_path(item["path"])
            else:
                attachment = None
        else:
            attachment = None

        if not attachment:
            continue

        key = (attachment.get("name"), attachment.get("size"))
        if key in seen:
            continue
        seen.add(key)
        result.append(attachment)
    return result


def images(attachments) -> list:
    return [item for item in attachments or [] if item.get("kind") == "image" and item.get("data")]


def image_payloads(attachments) -> list:
    return [item["data"] for item in images(attachments)]


def data_url(attachment) -> str:
    mime = attachment.get("mime") or "image/png"
    return f"data:{mime};base64,{attachment.get('data', '')}"


def describe(attachments) -> str:
    """文档正文拼成一段上下文；图片不用（模型看得见）。"""
    blocks = []
    for item in attachments or []:
        text = (item.get("text") or "").strip()
        if not text:
            continue

        head = f"【附件：{item.get('name')}（{human_size(item.get('size'))}）】"
        if item.get("truncated"):
            head += "（只取了前一部分）"
        blocks.append(f"{head}\n{text}")
    return "\n\n".join(blocks)


def with_document_context(text, attachments) -> str:
    context = describe(attachments)
    text = str(text or "")
    if not context:
        return text
    return f"{context}\n\n{text}" if text.strip() else context


def summarize(attachments) -> str:
    """给提示/日志用的一行描述。"""
    parts = []
    for item in attachments or []:
        mark = "图片" if item.get("kind") == "image" else "文件"
        parts.append(f"{mark} {item.get('name')}（{human_size(item.get('size'))}）")
    return "、".join(parts)


def ollama_message(text, attachments) -> dict:
    message = {"role": "user", "content": text}
    payloads = image_payloads(attachments)
    if payloads:
        message["images"] = payloads
    return message


def openai_message(text, attachments) -> dict:
    payloads = images(attachments)
    if not payloads:
        return {"role": "user", "content": text}

    content = []
    if str(text or "").strip():
        content.append({"type": "text", "text": text})
    for item in payloads:
        content.append({"type": "image_url", "image_url": {"url": data_url(item)}})
    return {"role": "user", "content": content}


def compact_for_display(value, keep=48):
    """记忆面板用的：把 base64 换成占位符，别让界面里塞几 MB 的字符串。"""
    def shorten(text: str) -> str:
        if len(text) <= keep * 2:
            return text
        return f"{text[:keep]}…<{human_size(len(text))} base64>…{text[-8:]}"

    if isinstance(value, str):
        return shorten(value)
    if isinstance(value, list):
        return [compact_for_display(item, keep) for item in value]
    if isinstance(value, dict):
        return {key: compact_for_display(item, keep) for key, item in value.items()}
    return value


_VISION_CACHE: dict = {}


def vision_capability(model: str, use_cache: bool = True):
    model = str(model or "").strip()
    if not model:
        return None
    if use_cache and model in _VISION_CACHE:
        return _VISION_CACHE[model]

    result = None
    try:
        import ollama

        info = ollama.show(model)
        capabilities = info.get("capabilities") if isinstance(info, dict) else getattr(info, "capabilities", None)
        if capabilities is not None:
            result = "vision" in [str(item).lower() for item in capabilities]
    except Exception:  # noqa: BLE001 - 查不到就当未知，不影响发消息
        result = None

    if result is not None:
        _VISION_CACHE[model] = result
    return result


def can_see_images(model: str) -> bool:
    """明确不支持看图才返回 False（未知按支持处理，避免误报）。"""
    return vision_capability(model) is not False
